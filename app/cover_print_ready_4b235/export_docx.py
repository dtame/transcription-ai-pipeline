"""One-page cover DOCX files. The background fills the bleed; the type is real text."""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt
from lxml import etree

from app.cover_print_ready_4b235.fonts import ensure_fonts, font_metrics
from app.cover_print_ready_4b235.guard import assert_output_allowed
from app.cover_print_ready_4b235.layout import DrawnLine, FaceLayout

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_V = "urn:schemas-microsoft-com:vml"
_O = "urn:schemas-microsoft-com:office:office"
_W10 = "urn:schemas-microsoft-com:office:word"


def render_front_docx(path: Path, layout: FaceLayout, image_path: Path, document_title: str) -> None:
    _render(path, layout, image_path, document_title)


def render_back_docx(path: Path, layout: FaceLayout, image_path: Path, document_title: str) -> None:
    _render(path, layout, image_path, document_title)


def _render(path: Path, layout: FaceLayout, image_path: Path, document_title: str) -> None:
    assert_output_allowed(path)
    fonts = ensure_fonts()
    family = str(fonts["family"])
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Emu(_inches_to_emu(layout.page_width_in))
    section.page_height = Emu(_inches_to_emu(layout.page_height_in))
    section.left_margin = Emu(0)
    section.right_margin = Emu(0)
    section.top_margin = Emu(0)
    section.bottom_margin = Emu(0)
    section.header_distance = Emu(0)
    section.footer_distance = Emu(0)
    section.gutter = Emu(0)
    _silence(section.header)
    _silence(section.footer)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    _add_bleed_picture(paragraph, image_path, layout.page_width_in, layout.page_height_in)
    if layout.lines:
        paragraph._p.append(_textbox(layout, family))
    document.core_properties.title = document_title
    document.core_properties.author = ""
    document.core_properties.comments = "Print-review cover. One page. 3 mm bleed. No page number."
    document.save(str(path))


def _add_bleed_picture(paragraph, image_path: Path, width_in: float, height_in: float) -> None:
    run = paragraph.add_run()
    run.add_picture(
        str(image_path),
        width=Emu(_inches_to_emu(width_in)),
        height=Emu(_inches_to_emu(height_in)),
    )
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "behindDoc": "1",
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
        "relativeHeight": "0",
    }.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    anchor.append(_offset("wp:positionH", "0"))
    anchor.append(_offset("wp:positionV", "0"))
    for tag in ("wp:extent", "wp:effectExtent"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(child)
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(child)
    drawing.remove(inline)
    drawing.append(anchor)


def _offset(tag: str, value: str):
    node = OxmlElement(tag)
    node.set("relativeFrom", "page")
    offset = OxmlElement("wp:posOffset")
    offset.text = value
    node.append(offset)
    return node


def _textbox(layout: FaceLayout, family: str):
    tops = [_line_top(line) for line in layout.lines]
    bottoms = [_line_bottom(line) for line in layout.lines]
    top_pt = max(0.0, min(tops) - 2)
    height_pt = max(bottoms) - top_pt + 8
    left_pt = layout.left_in * 72
    width_pt = layout.content_width_pt
    style = (
        "position:absolute;"
        f"margin-left:{left_pt:.2f}pt;margin-top:{top_pt:.2f}pt;"
        f"width:{width_pt:.2f}pt;height:{height_pt:.2f}pt;"
        "z-index:251659264;visibility:visible;mso-wrap-style:square;"
        "mso-position-horizontal:absolute;mso-position-horizontal-relative:page;"
        "mso-position-vertical:absolute;mso-position-vertical-relative:page"
    )
    paragraphs = "".join(_paragraph_xml(line, family, layout.lines, index) for index, line in enumerate(layout.lines))
    xml = (
        f'<w:r xmlns:w="{_W}" xmlns:v="{_V}" xmlns:o="{_O}" xmlns:w10="{_W10}">'
        "<w:pict>"
        f'<v:shape id="CoverText" type="#_x0000_t202" filled="f" stroked="f" style="{style}">'
        '<v:textbox inset="0,0,0,0">'
        f"<w:txbxContent>{paragraphs}</w:txbxContent>"
        "</v:textbox>"
        '<w10:wrap type="none"/>'
        "</v:shape>"
        "</w:pict>"
        "</w:r>"
    )
    return etree.fromstring(xml)


def _paragraph_xml(line: DrawnLine, family: str, lines: tuple[DrawnLine, ...], index: int) -> str:
    if index + 1 < len(lines):
        spacing = max(line.size_pt, lines[index + 1].baseline_from_top_pt - line.baseline_from_top_pt)
    else:
        spacing = line.size_pt * 1.15
    before = 0
    if index == 0:
        before = 0
    align = "center" if line.align == "center" else "left"
    bold = "<w:b/>" if line.kind == "bold" else ""
    italic = "<w:i/>" if line.kind == "italic" else ""
    color = "{:02X}{:02X}{:02X}".format(*line.color)
    half_points = int(round(line.size_pt * 2))
    tracking = int(round(line.tracking_pt * 20))
    return (
        "<w:p>"
        "<w:pPr>"
        f'<w:spacing w:before="{int(before)}" w:after="0" '
        f'w:line="{int(round(spacing * 20))}" w:lineRule="exact"/>'
        f'<w:jc w:val="{align}"/>'
        "</w:pPr>"
        "<w:r><w:rPr>"
        f'<w:rFonts w:ascii="{family}" w:hAnsi="{family}" w:cs="{family}"/>'
        f"{bold}{italic}"
        f'<w:sz w:val="{half_points}"/><w:szCs w:val="{half_points}"/>'
        f'<w:color w:val="{color}"/>'
        f'<w:spacing w:val="{tracking}"/>'
        "</w:rPr>"
        f'<w:t xml:space="preserve">{escape(line.text)}</w:t>'
        "</w:r></w:p>"
    )


def _line_top(line: DrawnLine) -> float:
    ascent, _descent = font_metrics(line.kind, line.size_pt)
    return line.baseline_from_top_pt - ascent


def _line_bottom(line: DrawnLine) -> float:
    _ascent, descent = font_metrics(line.kind, line.size_pt)
    return line.baseline_from_top_pt + descent


def _silence(part) -> None:
    for paragraph in part.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.text = ""


def _inches_to_emu(value: float) -> int:
    return int(round(value * 914400))


__all__ = ["render_back_docx", "render_front_docx"]
