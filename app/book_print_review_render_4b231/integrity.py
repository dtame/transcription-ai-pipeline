"""Re-open the generated DOCX and compare it to book.json."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT

from app.book_print_review_render_4b231.constants import (
    BOOK_TITLE,
    EXPECTED_CHAPTER_COUNT,
    EXPECTED_PARAGRAPH_COUNT,
    EXPECTED_SECTION_COUNT,
    PHASE,
)
from app.word_renderer.constants import (
    DRAFT_NOTICE,
    STYLE_BODY,
    STYLE_BODY_FIRST,
    STYLE_CHAPTER_NUMBER,
    STYLE_CHAPTER_TITLE,
    STYLE_SECTION_TITLE,
    TECHNICAL_STATUS_TOKENS,
)
from app.word_renderer.front_matter import version_notice
from app.word_renderer.geometry import inspect_geometry
from app.word_renderer.mapping import PrintBook, find_technical_leaks
from app.word_renderer.oxml import field_instructions
from app.word_renderer.styles import REQUIRED_STYLES


def collect_fields(doc) -> tuple[str, ...]:
    found: list[str] = []
    found.extend(field_instructions(doc.element))
    for section in doc.sections:
        for part in (
            section.header,
            section.footer,
            section.even_page_header,
            section.even_page_footer,
            section.first_page_header,
            section.first_page_footer,
        ):
            found.extend(field_instructions(part._element))
    return tuple(found)


def open_document(data: bytes):
    return Document(BytesIO(data))


def editorial_paragraphs(doc) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, paragraph in enumerate(doc.paragraphs):
        style = paragraph.style.name if paragraph.style is not None else ""
        if style not in {STYLE_BODY, STYLE_BODY_FIRST}:
            continue
        rows.append(
            {
                "word_index": index,
                "style": style,
                "text": paragraph.text,
            }
        )
    return rows


def structured_titles(doc) -> dict[str, list[str]]:
    chapters: list[str] = []
    sections: list[str] = []
    numbers: list[str] = []
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        text = paragraph.text.strip()
        if not text:
            continue
        if style == STYLE_CHAPTER_TITLE:
            chapters.append(text)
        elif style == STYLE_SECTION_TITLE:
            sections.append(text)
        elif style == STYLE_CHAPTER_NUMBER:
            numbers.append(text)
    return {"chapters": chapters, "sections": sections, "numbers": numbers}


def content_mapping(doc, book: PrintBook) -> dict[str, Any]:
    editorial = editorial_paragraphs(doc)
    canonical = list(book.paragraph_texts)
    rows: list[dict[str, Any]] = []
    cursor = 0
    for chapter in book.chapters:
        for section in chapter.sections:
            for paragraph in section.paragraphs:
                word = editorial[cursor] if cursor < len(editorial) else None
                rows.append(
                    {
                        "canonical_index": cursor,
                        "chapter_order": chapter.order,
                        "chapter_title": chapter.title,
                        "section_order": section.order,
                        "section_title": section.title,
                        "canonical_text": paragraph.text,
                        "word_index": None if word is None else word["word_index"],
                        "word_style": None if word is None else word["style"],
                        "word_text": None if word is None else word["text"],
                        "match": bool(word and word["text"] == paragraph.text),
                    }
                )
                cursor += 1
    extra = editorial[cursor:]
    return {
        "phase": PHASE,
        "canonical_count": len(canonical),
        "word_editorial_count": len(editorial),
        "mapped": rows,
        "extra_word_paragraphs": extra,
        "all_match": (
            len(editorial) == len(canonical)
            and all(row["match"] for row in rows)
            and not extra
        ),
        "verification": "automatic",
        "secrets_included": False,
    }


def validate_docx_integrity(
    data: bytes,
    book: PrintBook,
    profile: dict[str, Any],
) -> dict[str, Any]:
    doc = open_document(data)
    geometry_rows = [inspect_geometry(section, profile) for section in doc.sections]
    titles = structured_titles(doc)
    mapping = content_mapping(doc, book)
    fields = collect_fields(doc)
    styles_present = all(name in [style.name for style in doc.styles] for name in REQUIRED_STYLES)
    printed = tuple(
        paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()
    )
    leaks = find_technical_leaks(printed)
    invented = _invented_strings(printed, book)
    expected_numbers = [str(chapter.order) for chapter in book.chapters]
    checks = {
        "section_count": len(doc.sections) == EXPECTED_CHAPTER_COUNT + 1,
        "page_width": all(row["page_width_matches"] for row in geometry_rows),
        "page_height": all(row["page_height_matches"] for row in geometry_rows),
        "portrait": all(
            row["portrait"] and section.orientation == WD_ORIENT.PORTRAIT
            for row, section in zip(geometry_rows, doc.sections, strict=False)
        ),
        "gutter": all(row["gutter_matches_profile"] for row in geometry_rows),
        "gutter_not_duplicated": all(row["gutter_not_added_to_left_margin"] for row in geometry_rows),
        "chapters": titles["chapters"] == [chapter.title for chapter in book.chapters],
        "sections": titles["sections"]
        == [section.title for chapter in book.chapters for section in chapter.sections],
        "chapter_numbers": titles["numbers"] == expected_numbers,
        "paragraphs": mapping["all_match"]
        and mapping["canonical_count"] == EXPECTED_PARAGRAPH_COUNT,
        "styles_present": styles_present,
        "toc_field": any(item.startswith("TOC") for item in fields),
        "page_field": any(item == "PAGE" or item.startswith("PAGE") for item in fields),
        "styleref_field": any(item.startswith("STYLEREF") for item in fields),
        "no_technical_ids": not leaks,
        "no_invented_content": not invented,
        "draft_notice_separated": DRAFT_NOTICE in printed,
        "core_author_empty": not (doc.core_properties.author or ""),
        "title_metadata": (doc.core_properties.title or "") == BOOK_TITLE,
    }
    status = "PASS" if all(checks.values()) else "FAIL"
    return {
        "phase": PHASE,
        "status": status,
        "chapter_count": len(titles["chapters"]),
        "section_count": len(titles["sections"]),
        "paragraph_count": mapping["word_editorial_count"],
        "expected_chapters": EXPECTED_CHAPTER_COUNT,
        "expected_sections": EXPECTED_SECTION_COUNT,
        "expected_paragraphs": EXPECTED_PARAGRAPH_COUNT,
        "geometry": geometry_rows,
        "fields": {
            "toc": [item for item in fields if item.startswith("TOC")],
            "page": [item for item in fields if item == "PAGE" or item.startswith("PAGE")],
            "styleref": [item for item in fields if item.startswith("STYLEREF")],
        },
        "technical_leaks": list(leaks),
        "invented_content": invented,
        "checks": checks,
        "mapping_all_match": mapping["all_match"],
        "verification": "automatic",
        "secrets_included": False,
        "mapping": mapping,
    }


def _invented_strings(printed: tuple[str, ...], book: PrintBook) -> list[str]:
    allowed = {
        book.title,
        DRAFT_NOTICE,
        "Contents",
        "Right-click to update the table of contents.",
    }
    if book.subtitle:
        allowed.add(book.subtitle)
    notice = version_notice(book)
    if notice:
        allowed.add(notice)
    allowed.update(chapter.title for chapter in book.chapters)
    allowed.update(
        section.title for chapter in book.chapters for section in chapter.sections
    )
    allowed.update(book.paragraph_texts)
    allowed.update(str(chapter.order) for chapter in book.chapters)
    invented: list[str] = []
    for text in printed:
        stripped = text.strip()
        if not stripped:
            continue
        if stripped in allowed:
            continue
        if stripped.isdigit():
            continue
        if any(token in stripped for token in TECHNICAL_STATUS_TOKENS):
            invented.append(stripped)
            continue
        invented.append(stripped)
    return invented


__all__ = [
    "collect_fields",
    "content_mapping",
    "editorial_paragraphs",
    "open_document",
    "validate_docx_integrity",
]
