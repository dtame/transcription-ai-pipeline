"""Phase 4B.2.31 — offline print-review DOCX/PDF generation."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches

from app.book_print_review_canonical_4b229.constants import (
    AUTHORIZATION_SCOPE as CONSUMED_4B229_SCOPE,
)
from app.word_print_profile_4b230.constants import (
    AUTHORIZATION_SCOPE as CONSUMED_4B230_SCOPE,
)
from app.word_print_profile_4b230.validation import synthetic_book_payload
from app.book_print_review_render_4b231.constants import (
    AUTHORIZATION_SCOPE,
    AUTHORIZED_ANTHROPIC_CALLS,
    AUTHORIZED_OPENAI_CALLS,
    AUTHORIZED_SONNET_CALLS,
    AUTHORIZED_TERRA_CALLS,
    BOOK_JSON_MUTATION_AUTHORIZED,
    BOOK_STATUS,
    BOOK_TITLE,
    BOOK_VERSION,
    COVER_GENERATION_AUTHORIZED,
    DOCX_GENERATION_AUTHORIZED,
    EXPECTED_BOOK_SHA256,
    EXPECTED_CHAPTER_COUNT,
    EXPECTED_PARAGRAPH_COUNT,
    EXPECTED_SECTION_COUNT,
    FINAL_PUBLICATION_AUTHORIZED,
    PDF_GENERATION_AUTHORIZED,
    PRINT_PROFILE,
    PUBLICATION_STATUS,
)
from app.book_print_review_render_4b231.generation import build_docx_bytes, write_docx_bytes
from app.book_print_review_render_4b231.guard import (
    BookPrintReviewRender4231Error,
    assert_offline_only,
    assert_publication_target_allowed,
    validate_authorization_scope,
)
from app.book_print_review_render_4b231.hashes import snapshot
from app.book_print_review_render_4b231.integrity import collect_fields, validate_docx_integrity
from app.book_print_review_render_4b231.paths import (
    official_publication_root,
    production_book_path,
)
from app.book_print_review_render_4b231.preflight import build_preflight
from app.book_print_review_render_4b231.publication import (
    assert_replace_allowed,
    inspect_existing_docx,
    publish_print_review,
)
from app.book_print_review_render_4b231.runner import run_phase
from app.word_renderer.constants import (
    DRAFT_NOTICE,
    GUTTER_INCHES,
    INSIDE_MARGIN_INCHES,
    PAGE_HEIGHT_INCHES,
    PAGE_WIDTH_INCHES,
    STYLE_BODY,
    STYLE_BODY_FIRST,
    STYLE_CHAPTER_TITLE,
    STYLE_SECTION_TITLE,
)
from app.word_renderer.document import build_print_document
from app.word_renderer.finalizer import (
    FakeWordBackend,
    UnavailableWordBackend,
    WordFinalizerError,
    detect_word_environment,
    finalize_word_document,
    finalizer_readiness,
    word_progid_present,
)
from app.word_renderer.front_matter import version_notice
from app.word_renderer.mapping import map_book
from app.word_renderer.oxml import gutter_emu, page_number_start
from app.word_renderer.profile import load_profile
from app.word_renderer.styles import inspect_style


@pytest.fixture(scope="module")
def preflight_data():
    return build_preflight()


@pytest.fixture(scope="module")
def real_docx_bytes(preflight_data):
    return build_docx_bytes(preflight_data["book"], preflight_data["profile"])


@pytest.fixture(scope="module")
def real_integrity(real_docx_bytes, preflight_data):
    return validate_docx_integrity(
        real_docx_bytes,
        preflight_data["book"],
        preflight_data["profile"],
    )


def test_offline_authorizations_and_no_provider_calls():
    assert AUTHORIZED_ANTHROPIC_CALLS == 0
    assert AUTHORIZED_OPENAI_CALLS == 0
    assert AUTHORIZED_SONNET_CALLS == 0
    assert AUTHORIZED_TERRA_CALLS == 0
    assert DOCX_GENERATION_AUTHORIZED is True
    assert PDF_GENERATION_AUTHORIZED is True
    assert COVER_GENERATION_AUTHORIZED is False
    assert BOOK_JSON_MUTATION_AUTHORIZED is False
    assert FINAL_PUBLICATION_AUTHORIZED is False
    assert_offline_only()
    validate_authorization_scope(AUTHORIZATION_SCOPE)
    with pytest.raises(BookPrintReviewRender4231Error):
        validate_authorization_scope(CONSUMED_4B229_SCOPE)
    with pytest.raises(BookPrintReviewRender4231Error):
        validate_authorization_scope(CONSUMED_4B230_SCOPE)
    with pytest.raises(BookPrintReviewRender4231Error):
        validate_authorization_scope("WRONG")


def test_load_canonical_and_sha256(preflight_data):
    path = production_book_path()
    assert path.is_file()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == EXPECTED_BOOK_SHA256
    assert preflight_data["title"] == BOOK_TITLE
    assert preflight_data["version"] == BOOK_VERSION
    assert preflight_data["editorial_status"] == BOOK_STATUS
    assert preflight_data["book_sha256"] == EXPECTED_BOOK_SHA256


def test_load_print_profile():
    profile = load_profile()
    assert profile["profile_id"] == PRINT_PROFILE
    assert profile["book_agnostic"] is True
    assert BOOK_TITLE not in __import__("json").dumps(profile)


def test_geometry_mirror_margins_gutter_and_styles(real_docx_bytes):
    doc = Document(__import__("io").BytesIO(real_docx_bytes))
    section = doc.sections[1]
    assert abs(section.page_width.inches - PAGE_WIDTH_INCHES) < 0.0001
    assert abs(section.page_height.inches - PAGE_HEIGHT_INCHES) < 0.0001
    assert section.orientation == WD_ORIENT.PORTRAIT
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    assert doc.settings.element.find(f"{ns}mirrorMargins") is not None
    assert doc.settings.element.find(f"{ns}updateFields") is not None
    assert abs(section.left_margin.inches - INSIDE_MARGIN_INCHES) < 0.0001
    assert gutter_emu(section) == int(Inches(GUTTER_INCHES))
    body = inspect_style(doc, STYLE_BODY)
    first = inspect_style(doc, STYLE_BODY_FIRST)
    assert body["font_name"] == "Georgia"
    assert body["font_size_pt"] == 11
    assert first["first_line_indent_inches"] == 0
    assert inspect_style(doc, STYLE_CHAPTER_TITLE)["keep_with_next"] is True
    assert inspect_style(doc, STYLE_SECTION_TITLE)["keep_with_next"] is True
    assert page_number_start(doc.sections[1]) == 1
    assert all(page_number_start(section) is None for section in doc.sections[2:])


def test_front_matter_fields_and_no_invented_content(real_docx_bytes, preflight_data, real_integrity):
    doc = Document(__import__("io").BytesIO(real_docx_bytes))
    texts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    assert texts[0] == BOOK_TITLE
    assert DRAFT_NOTICE in texts
    assert version_notice(preflight_data["book"]) in texts
    assert "Contents" in texts
    assert BOOK_STATUS not in "\n".join(texts)
    fields = collect_fields(doc)
    assert any(item.startswith("TOC") for item in fields)
    assert any(item.startswith("STYLEREF") for item in fields)
    assert any(item == "PAGE" or item.startswith("PAGE") for item in fields)
    assert real_integrity["checks"]["no_invented_content"] is True
    assert real_integrity["status"] == "PASS"


def test_chapter_section_paragraph_identity(real_integrity, preflight_data):
    assert real_integrity["chapter_count"] == EXPECTED_CHAPTER_COUNT
    assert real_integrity["section_count"] == EXPECTED_SECTION_COUNT
    assert real_integrity["paragraph_count"] == EXPECTED_PARAGRAPH_COUNT
    assert real_integrity["mapping_all_match"] is True
    assert preflight_data["book"].paragraph_texts == tuple(
        preflight_data["book"].paragraph_texts
    )
    assert [chapter.order for chapter in preflight_data["book"].chapters] == list(range(1, 20))


def test_word_detection_does_not_assume_windows_is_enough():
    environment = detect_word_environment(probe=False, backend=UnavailableWordBackend())
    assert environment["word_available"] is False
    assert environment["assumed_from_windows_only"] is False
    assert environment["word_progid_present"] == word_progid_present()


def test_word_absent_keeps_docx_and_refuses_pdf(tmp_path, real_docx_bytes):
    source = tmp_path / "source.docx"
    working = tmp_path / "working.docx"
    pdf = tmp_path / "The_Life_You_Already_Inherited_print_review_v1.pdf"
    source.write_bytes(real_docx_bytes)
    result = finalize_word_document(
        source_docx=source,
        working_docx=working,
        pdf_path=pdf,
        backend=UnavailableWordBackend(),
    )
    assert result["status"] == "NOT_AVAILABLE"
    assert result["executed"] is False
    assert working.is_file()
    assert not pdf.exists()
    assert "Open the unfinalized DOCX" in result["required_human_action"]


def test_com_error_handling_does_not_claim_success(tmp_path, real_docx_bytes):
    source = tmp_path / "source.docx"
    working = tmp_path / "working.docx"
    source.write_bytes(real_docx_bytes)
    backend = FakeWordBackend(fail=True, error="RPC_E_SERVERFAULT")
    result = finalize_word_document(
        source_docx=source,
        working_docx=working,
        pdf_path=tmp_path / "out.pdf",
        backend=backend,
    )
    assert result["status"] == "FAIL"
    assert "RPC_E_SERVERFAULT" in result["error"]
    assert result["pdf_exported"] is False
    assert not (tmp_path / "out.pdf").exists()


def test_fake_word_updates_fields_and_exports_pdf(tmp_path, real_docx_bytes):
    source = tmp_path / "source.docx"
    working = tmp_path / "working.docx"
    pdf = tmp_path / "out.pdf"
    source.write_bytes(real_docx_bytes)
    result = finalize_word_document(
        source_docx=source,
        working_docx=working,
        pdf_path=pdf,
        backend=FakeWordBackend(),
    )
    assert result["status"] == "PASS"
    assert result["fields_updated"] is True
    assert result["toc_updated"] is True
    assert result["pdf_exported"] is True
    assert pdf.is_file()
    with pytest.raises(WordFinalizerError):
        finalize_word_document(
            source_docx=source,
            working_docx=source,
            backend=FakeWordBackend(),
        )


def test_field_update_and_pdf_export_when_word_available(tmp_path):
    environment = detect_word_environment(probe=False)
    if not environment.get("word_progid_present"):
        pytest.skip("Microsoft Word is not registered on this machine")
    profile = load_profile()
    book = map_book(synthetic_book_payload())
    source = tmp_path / "synthetic.docx"
    working = tmp_path / "synthetic_working.docx"
    pdf = tmp_path / "synthetic.pdf"
    write_docx_bytes(source, build_docx_bytes(book, profile))
    result = finalize_word_document(
        source_docx=source,
        working_docx=working,
        pdf_path=pdf,
        max_iterations=2,
    )
    assert result["status"] in {"PASS", "PARTIAL", "FAIL", "NOT_AVAILABLE"}
    if result["status"] == "NOT_AVAILABLE":
        assert not pdf.exists()
        return
    assert result["executed"] is True
    assert result["fields_updated"] is True
    if result["pdf_exported"]:
        assert pdf.is_file() and pdf.stat().st_size > 0
        assert pdf.read_bytes()[:4] == b"%PDF"
    assert official_publication_root() not in pdf.parents
    assert official_publication_root() not in working.parents


def test_overwrite_protection_and_atomic_publication(tmp_path, real_docx_bytes):
    source = tmp_path / "final.docx"
    source.write_bytes(real_docx_bytes)
    dest = tmp_path / "publication" / "book.docx"
    pdf_dest = tmp_path / "publication" / "book.pdf"
    first = publish_print_review(
        finalized_docx=source,
        pdf_path=None,
        destination_docx=dest,
        destination_pdf=pdf_dest,
        allow_official=False,
    )
    assert first["action"] == "create"
    assert dest.is_file()
    assert not pdf_dest.exists()
    assert not dest.with_name(dest.name + ".partial").exists()
    second = publish_print_review(
        finalized_docx=source,
        pdf_path=None,
        destination_docx=dest,
        destination_pdf=pdf_dest,
        allow_official=False,
    )
    assert second["action"] == "idempotent"
    protected = tmp_path / "protected.docx"
    write_docx_bytes(protected, real_docx_bytes)
    doc = Document(protected)
    doc.core_properties.category = "COMMERCIAL"
    doc.save(protected)
    existing = inspect_existing_docx(protected)
    assert existing["protected"] is True
    with pytest.raises(BookPrintReviewRender4231Error):
        assert_replace_allowed(existing, new_sha256="different")
    with pytest.raises(BookPrintReviewRender4231Error):
        assert_publication_target_allowed(
            official_publication_root() / "probe.docx",
            allow_official=False,
        )


def test_sources_unchanged_and_no_official_pollution(tmp_path):
    before = snapshot()
    result = run_phase(
        authorization_scope=AUTHORIZATION_SCOPE,
        write_artifacts=True,
        run_tests=False,
        finalize=True,
        probe_word=False,
        allow_official=False,
        root=tmp_path,
        word_backend=UnavailableWordBackend(),
    )
    after = snapshot()
    assert before["book_json"] == after["book_json"]
    assert after["book_json"]["sha256"] == EXPECTED_BOOK_SHA256
    assert result.bundle["header"]["provider_calls"] == 0
    assert result.bundle["header"]["publication_status"] == PUBLICATION_STATUS
    assert result.bundle["header"]["docx_content_integrity"] == "PASS"
    assert result.mode in {"PARTIAL", "PASS"}
    official = official_publication_root()
    published = list((tmp_path / "sortie").rglob("*.docx"))
    assert published
    assert all(official.resolve() not in path.resolve().parents for path in published)
    assert not any(path.suffix == ".pdf" for path in (tmp_path / "sortie").rglob("*"))


def test_closed_readiness_contract_unchanged():
    payload = finalizer_readiness()
    assert payload["executed"] is False
    assert payload["status"] == "PARTIAL"
    assert payload["exists_in_repository"] is False


def test_build_print_document_still_reusable_for_other_books():
    profile = load_profile()
    other = map_book(synthetic_book_payload())
    doc = build_print_document(other, profile)
    assert doc.core_properties.title == "A Different Synthetic Book"
    assert version_notice(other) == "Version synthetic-v1"
