"""Configure in-memory Word documents. Never publishes a DOCX or PDF."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.enum.section import WD_SECTION

from app.word_renderer.constants import STYLE_BODY, STYLE_BODY_FIRST, STYLE_CHAPTER_NUMBER, STYLE_CHAPTER_TITLE, STYLE_SECTION_TITLE
from app.word_renderer.front_matter import add_front_matter
from app.word_renderer.geometry import (
    apply_body_section,
    apply_document_settings,
    apply_front_matter_section,
    chapter_start_type,
)
from app.word_renderer.headers import apply_body_headers, apply_front_matter_headers
from app.word_renderer.mapping import PrintBook
from app.word_renderer.oxml import count_explicit_page_breaks, field_instructions
from app.word_renderer.profile import chapter_start_type_name, load_profile
from app.word_renderer.styles import apply_styles

PUBLISH_FORBIDDEN_SUFFIXES = {".docx", ".pdf"}


class WordPublishForbidden(RuntimeError):
    """4B.2.30 must not write a published DOCX or PDF."""


def configure_document(profile: Mapping[str, Any] | None = None):
    resolved = dict(profile or load_profile())
    doc = Document()
    apply_document_settings(doc, resolved)
    apply_styles(doc, resolved)
    apply_front_matter_section(doc.sections[0], resolved)
    apply_front_matter_headers(doc.sections[0], resolved)
    if doc.core_properties.author:
        doc.core_properties.author = ""
    return doc


def add_body_chapter(
    doc,
    chapter,
    profile: Mapping[str, Any],
    *,
    book_title: str,
    first: bool,
) -> None:
    if first:
        body = doc.add_section(chapter_start_type(profile, first=True))
        apply_body_section(body, profile, start=True)
        apply_body_headers(body, book_title, profile)
    else:
        body = doc.add_section(chapter_start_type(profile, first=False))
        apply_body_section(body, profile, start=False)
        for part in (
            body.header,
            body.footer,
            body.even_page_header,
            body.even_page_footer,
            body.first_page_header,
            body.first_page_footer,
        ):
            part.is_linked_to_previous = True
    numbering = (profile.get("chapter") or {}).get("numbering") or {}
    if numbering.get("display") == "separate_line":
        template = str(numbering.get("label_template") or "{order}")
        doc.add_paragraph(template.format(order=chapter.order), style=STYLE_CHAPTER_NUMBER)
    title = chapter.title
    if numbering.get("include_in_title_text"):
        raise WordPublishForbidden("refusing to merge chapter numbers into titles")
    doc.add_paragraph(title, style=STYLE_CHAPTER_TITLE)
    for section in chapter.sections:
        doc.add_paragraph(section.title, style=STYLE_SECTION_TITLE)
        for paragraph in section.paragraphs:
            style = STYLE_BODY_FIRST if paragraph.is_first else STYLE_BODY
            doc.add_paragraph(paragraph.text, style=style)


def build_print_document(
    book: PrintBook,
    profile: Mapping[str, Any] | None = None,
    *,
    include_sections: bool | None = None,
):
    resolved = dict(profile or load_profile())
    doc = configure_document(resolved)
    doc.core_properties.title = book.title
    doc.core_properties.author = ""
    add_front_matter(doc, book, resolved, include_sections=include_sections)
    for index, chapter in enumerate(book.chapters):
        add_body_chapter(
            doc,
            chapter,
            resolved,
            book_title=book.title,
            first=index == 0,
        )
    return doc


def serialize_in_memory(doc) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def assert_not_published(path: Path) -> None:
    if path.suffix.lower() in PUBLISH_FORBIDDEN_SUFFIXES:
        raise WordPublishForbidden(
            f"writing {path.suffix} is forbidden during print-profile preparation"
        )


def inspect_document(doc, profile: Mapping[str, Any]) -> dict[str, Any]:
    body_starts = [
        str(section.start_type)
        for section in doc.sections[1:]
    ]
    interchapter = body_starts[1:]
    new_page = str(WD_SECTION.NEW_PAGE)
    odd_page = str(WD_SECTION.ODD_PAGE)
    page_breaks = count_explicit_page_breaks(doc.element)
    return {
        "section_count": len(doc.sections),
        "front_matter_start_type": str(doc.sections[0].start_type) if doc.sections else "",
        "body_start_types": body_starts,
        "interchapter_start_types": interchapter,
        "chapter_start_policy": chapter_start_type_name(profile),
        "odd_or_new_page_starts": all(
            start in {odd_page, new_page} for start in body_starts
        ),
        "interchapter_next_page": all(start == new_page for start in interchapter),
        "interchapter_odd_page": any(start == odd_page for start in interchapter),
        "explicit_page_breaks": page_breaks,
        "toc_fields": tuple(
            item for item in field_instructions(doc.element) if item.startswith("TOC")
        ),
        "styleref_fields": tuple(
            item for item in field_instructions(doc.element) if item.startswith("STYLEREF")
        ),
        "page_fields": tuple(
            item for item in field_instructions(doc.element) if item == "PAGE" or item.startswith("PAGE")
        ),
        "core_author": doc.core_properties.author or "",
        "core_title": doc.core_properties.title or "",
    }


__all__ = [
    "WordPublishForbidden",
    "add_body_chapter",
    "assert_not_published",
    "build_print_document",
    "configure_document",
    "inspect_document",
    "serialize_in_memory",
]
