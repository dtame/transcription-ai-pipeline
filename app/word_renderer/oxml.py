"""OOXML helpers for print layout. Reuses the field pattern from docx_export_service."""

from __future__ import annotations

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Twips


def _child(parent, tag: str):
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    return element


def set_boolean_setting(settings, tag: str, enabled: bool) -> None:
    element = _child(settings, tag)
    if enabled:
        if qn("w:val") in element.attrib:
            del element.attrib[qn("w:val")]
    else:
        element.set(qn("w:val"), "0")


def enable_mirror_margins(doc) -> None:
    set_boolean_setting(doc.settings.element, "w:mirrorMargins", True)


def enable_even_odd_headers(doc) -> None:
    set_boolean_setting(doc.settings.element, "w:evenAndOddHeaders", True)


def disable_hyphenation(doc) -> None:
    set_boolean_setting(doc.settings.element, "w:autoHyphenation", False)


def enable_update_fields_on_open(doc) -> None:
    """Ask Word to refresh fields when a human opens the file.

    This is a convenience for unfinalized drafts. It is not a substitute
    for a real Word field-update pass.
    """
    set_boolean_setting(doc.settings.element, "w:updateFields", True)


def set_gutter(section, inches: float) -> None:
    if hasattr(section, "gutter"):
        section.gutter = Inches(inches)
        return
    pg_mar = _child(section._sectPr, "w:pgMar")
    pg_mar.set(qn("w:gutter"), str(int(Twips(inches * 1440))))


def gutter_emu(section) -> int:
    if hasattr(section, "gutter") and section.gutter is not None:
        return int(section.gutter)
    pg_mar = section._sectPr.find(qn("w:pgMar"))
    if pg_mar is None:
        return 0
    return int(pg_mar.get(qn("w:gutter") or "") or 0)


def set_page_number_format(section, *, fmt: str, start: int | None = None) -> None:
    element = _child(section._sectPr, "w:pgNumType")
    if fmt == "none":
        element.set(qn("w:fmt"), "none")
    else:
        element.set(qn("w:fmt"), fmt)
    if start is not None:
        element.set(qn("w:start"), str(start))
    elif qn("w:start") in element.attrib:
        del element.attrib[qn("w:start")]


def page_number_start(section) -> int | None:
    element = section._sectPr.find(qn("w:pgNumType"))
    if element is None:
        return None
    value = element.get(qn("w:start"))
    return int(value) if value is not None else None


def set_outline_level(style, level: int | None) -> None:
    p_pr = style.element.get_or_add_pPr()
    existing = p_pr.find(qn("w:outlineLvl"))
    if level is None:
        if existing is not None:
            p_pr.remove(existing)
        return
    if existing is None:
        existing = OxmlElement("w:outlineLvl")
        p_pr.append(existing)
    existing.set(qn("w:val"), str(int(level)))


def outline_level(style) -> int | None:
    p_pr = style.element.pPr
    if p_pr is None:
        return None
    existing = p_pr.find(qn("w:outlineLvl"))
    if existing is None:
        return None
    value = existing.get(qn("w:val"))
    return int(value) if value is not None else None


def set_rfonts(style, font_name: str) -> None:
    r_pr = style.element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.append(fonts)
    fonts.set(qn("w:ascii"), font_name)
    fonts.set(qn("w:hAnsi"), font_name)
    fonts.set(qn("w:cs"), font_name)


def suppress_hyphenation(style) -> None:
    p_pr = style.element.get_or_add_pPr()
    _child(p_pr, "w:suppressAutoHyphens")


def add_complex_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    instr_run._r.append(instr)

    sep_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(separate)

    if placeholder:
        paragraph.add_run(placeholder)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_page_field(paragraph) -> None:
    add_complex_field(paragraph, "PAGE")


def add_styleref_field(paragraph, style_name: str) -> None:
    add_complex_field(paragraph, f'STYLEREF "{style_name}"')


def add_toc_field(paragraph, *, include_sections: bool) -> None:
    levels = "1-2" if include_sections else "1-1"
    add_complex_field(
        paragraph,
        f'TOC \\o "{levels}" \\h \\z \\u',
        "Right-click to update the table of contents.",
    )


def field_instructions(container) -> tuple[str, ...]:
    texts: list[str] = []
    for node in container.iter(qn("w:instrText")):
        text = (node.text or "").strip()
        if text:
            texts.append(text)
    return tuple(texts)


def count_explicit_page_breaks(container) -> int:
    return sum(
        1
        for node in container.iter(qn("w:br"))
        if node.get(qn("w:type")) == "page"
    )


def iter_paragraph_texts(container) -> tuple[str, ...]:
    texts: list[str] = []
    for node in container.iter(qn("w:t")):
        text = node.text or ""
        if text:
            texts.append(text)
    return tuple(texts)


__all__ = [
    "add_page_field",
    "add_styleref_field",
    "add_toc_field",
    "count_explicit_page_breaks",
    "disable_hyphenation",
    "enable_even_odd_headers",
    "enable_mirror_margins",
    "enable_update_fields_on_open",
    "field_instructions",
    "gutter_emu",
    "iter_paragraph_texts",
    "outline_level",
    "page_number_start",
    "set_gutter",
    "set_outline_level",
    "set_page_number_format",
    "set_rfonts",
    "suppress_hyphenation",
]
