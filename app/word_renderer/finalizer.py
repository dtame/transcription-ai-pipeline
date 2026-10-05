"""Word Finalizer.

4B.2.30 used finalizer_readiness() as a contract that does not execute.
4B.2.31 implements the real Microsoft Word pass: field refresh, TOC
refresh, pagination, and PDF export.

The readiness payload stays unchanged so the closed 4B.2.30 phase
remains historically accurate. New callers use detect_word_environment
and finalize_word_document.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Callable, Mapping, Protocol


class WordFinalizerError(RuntimeError):
    """Microsoft Word finalization failed or is unavailable."""


class WordBackend(Protocol):
    name: str

    def detect(self) -> dict[str, Any]:
        ...

    def finalize(
        self,
        *,
        working_docx: Path,
        pdf_path: Path | None,
        max_iterations: int,
    ) -> dict[str, Any]:
        ...


def finalizer_readiness() -> dict[str, Any]:
    return {
        "component": "Word Finalizer",
        "exists_in_repository": False,
        "status": "PARTIAL",
        "executed": False,
        "prepared_contract": True,
        "requires_microsoft_word_or_equivalent": True,
        "python_docx_can": [
            "write_section_geometry",
            "write_styles",
            "write_toc_field",
            "write_styleref_field",
            "write_page_field",
            "write_odd_page_section_breaks",
            "write_mirror_margins_and_gutter",
        ],
        "python_docx_cannot": [
            "compute_final_page_numbers",
            "refresh_toc_entries",
            "materialize_blank_pages_for_odd_starts",
            "stabilize_pagination_after_reflow",
            "export_print_accurate_pdf",
        ],
        "planned_operations": [
            {
                "id": "open_docx",
                "description": "Open the generated DOCX in Microsoft Word.",
                "requires_real_renderer": True,
            },
            {
                "id": "update_fields",
                "description": "Update PAGE, STYLEREF, and other fields.",
                "requires_real_renderer": True,
            },
            {
                "id": "update_toc",
                "description": "Refresh the TOC field after pagination settles.",
                "requires_real_renderer": True,
            },
            {
                "id": "stabilize_pagination",
                "description": "Repaginate until page numbers stop changing.",
                "requires_real_renderer": True,
            },
            {
                "id": "verify_sections",
                "description": "Confirm 6x9 geometry, mirror margins, and chapter starts.",
                "requires_real_renderer": True,
            },
            {
                "id": "export_pdf",
                "description": "Export a print PDF from the refreshed DOCX.",
                "requires_real_renderer": True,
                "authorized_in_this_phase": False,
            },
        ],
        "windows_automation_reserved": "win32com / Word COM is reserved for Phase 4B.2.31+",
        "pagination_validated": False,
        "pdf_export_authorized": False,
        "notes": (
            "This phase only prepares the contract. python-docx writes fields "
            "but cannot refresh them. Page numbers shown only after a real "
            "Word pass must not be treated as validated here."
        ),
    }


def word_progid_present() -> bool:
    try:
        import winreg
    except ImportError:
        return False
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application"):
            return True
    except OSError:
        return False


def winword_process_ids() -> tuple[int, ...]:
    try:
        completed = subprocess.run(
            ["tasklist", "/FI", "IMAGENAME eq WINWORD.EXE", "/FO", "CSV", "/NH"],
            capture_output=True,
            text=True,
            check=False,
            timeout=15,
        )
    except (OSError, subprocess.TimeoutExpired):
        return ()
    pids: list[int] = []
    for line in (completed.stdout or "").splitlines():
        if "WINWORD.EXE" not in line.upper():
            continue
        parts = [part.strip().strip('"') for part in line.split(",")]
        if len(parts) >= 2 and parts[1].isdigit():
            pids.append(int(parts[1]))
    return tuple(pids)


def win32com_available() -> bool:
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        return False
    return True


def powershell_available() -> bool:
    return shutil.which("powershell.exe") is not None or shutil.which("powershell") is not None


def default_backend() -> WordBackend:
    if win32com_available():
        return Win32WordBackend()
    return PowerShellWordBackend()


def detect_word_environment(
    *,
    probe: bool = False,
    backend: WordBackend | None = None,
    probe_factory: Callable[[Path], None] | None = None,
) -> dict[str, Any]:
    chosen = backend or default_backend()
    detection = {
        "word_progid_present": word_progid_present(),
        "winword_process_running": bool(winword_process_ids()),
        "win32com_available": win32com_available(),
        "powershell_available": powershell_available(),
        "backend": chosen.name,
        "windows": os.name == "nt",
        "word_available": False,
        "assumed_from_windows_only": False,
        "probe_executed": False,
        "capabilities": {
            "open_save": "NOT_PROBED",
            "update_fields": "NOT_PROBED",
            "update_toc": "NOT_PROBED",
            "export_pdf": "NOT_PROBED",
        },
        "error": "",
        "verification": "automatic",
        "secrets_included": False,
    }
    if not detection["word_progid_present"]:
        detection["error"] = "Word.Application ProgID is not registered."
        return detection
    try:
        backend_detect = chosen.detect()
    except Exception as exc:
        detection["error"] = str(exc)
        return detection
    detection.update({key: value for key, value in backend_detect.items() if key != "capabilities"})
    if backend_detect.get("capabilities"):
        detection["capabilities"].update(backend_detect["capabilities"])
    detection["word_available"] = bool(
        detection["word_progid_present"] and backend_detect.get("backend_ready")
    )
    if not detection["word_available"]:
        detection["error"] = detection.get("error") or backend_detect.get("error") or (
            "A Word COM backend is not ready."
        )
        return detection
    if not probe:
        return detection
    detection["probe_executed"] = True
    try:
        probe_result = _probe_backend(chosen, probe_factory=probe_factory)
    except Exception as exc:
        detection["word_available"] = False
        detection["error"] = str(exc)
        detection["capabilities"] = {
            "open_save": "FAIL",
            "update_fields": "FAIL",
            "update_toc": "FAIL",
            "export_pdf": "FAIL",
        }
        return detection
    detection["capabilities"] = probe_result["capabilities"]
    detection["probe"] = probe_result
    if not probe_result.get("ok"):
        detection["word_available"] = False
        detection["error"] = probe_result.get("error") or "Word probe failed."
    return detection


def finalize_word_document(
    *,
    source_docx: Path,
    working_docx: Path,
    pdf_path: Path | None = None,
    backend: WordBackend | None = None,
    max_iterations: int = 3,
) -> dict[str, Any]:
    source = Path(source_docx)
    working = Path(working_docx)
    if not source.is_file():
        raise WordFinalizerError(f"source DOCX is missing: {source}")
    if source.resolve() == working.resolve():
        raise WordFinalizerError(
            "refusing to finalize the original unfinalized DOCX in place"
        )
    working.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, working)
    environment = detect_word_environment(probe=False, backend=backend)
    if not environment.get("word_available"):
        if pdf_path is not None and Path(pdf_path).exists():
            raise WordFinalizerError("refusing to keep a PDF when Word is unavailable")
        return {
            "status": "NOT_AVAILABLE",
            "executed": False,
            "fields_updated": False,
            "toc_updated": False,
            "pagination_stable": False,
            "pdf_exported": False,
            "docx_saved": False,
            "working_docx": str(working).replace("\\", "/"),
            "pdf_path": "",
            "environment": environment,
            "error": environment.get("error") or "Microsoft Word is not available.",
            "required_human_action": (
                "Open the unfinalized DOCX in Microsoft Word, update fields "
                "and the table of contents, stabilize pagination, export a "
                "print PDF, then re-run verification. Page numbers in the "
                "unfinalized DOCX are not definitive."
            ),
            "quit_called": False,
            "verification": "not_verified",
            "secrets_included": False,
        }
    chosen = backend or default_backend()
    try:
        raw = chosen.finalize(
            working_docx=working,
            pdf_path=Path(pdf_path) if pdf_path is not None else None,
            max_iterations=max_iterations,
        )
    except Exception as exc:
        return _failure_report(
            working=working,
            pdf_path=pdf_path,
            environment=environment,
            error=str(exc),
        )
    pdf_ok = False
    if pdf_path is not None:
        target = Path(pdf_path)
        pdf_ok = target.is_file() and target.stat().st_size > 0 and bool(raw.get("pdf_exported"))
        if not pdf_ok and target.exists() and target.stat().st_size == 0:
            target.unlink()
    status = str(raw.get("status") or "FAIL")
    if raw.get("ok") and raw.get("fields_updated") and raw.get("toc_updated"):
        if pdf_path is None or pdf_ok:
            status = "PASS" if raw.get("pagination_stable") else "PARTIAL"
        else:
            status = "PARTIAL"
    elif raw.get("ok"):
        status = "PARTIAL"
    else:
        status = "FAIL"
    return {
        "status": status,
        "executed": True,
        "backend": chosen.name,
        "fields_updated": bool(raw.get("fields_updated")),
        "toc_updated": bool(raw.get("toc_updated")),
        "pagination_stable": bool(raw.get("pagination_stable")),
        "stability_iterations": int(raw.get("stability_iterations") or 0),
        "page_count": int(raw.get("page_count") or 0),
        "toc_entries": list(raw.get("toc_entries") or []),
        "chapter_starts": list(raw.get("chapter_starts") or []),
        "pdf_exported": pdf_ok,
        "docx_saved": working.is_file() and bool(raw.get("docx_saved")),
        "working_docx": str(working).replace("\\", "/"),
        "pdf_path": str(pdf_path).replace("\\", "/") if pdf_ok and pdf_path is not None else "",
        "word_was_already_running": bool(raw.get("word_was_already_running")),
        "created_new_word_instance": bool(raw.get("created_new_word_instance")),
        "quit_called": bool(raw.get("quit_called")),
        "documents_closed_by_us": int(raw.get("documents_closed_by_us") or 0),
        "notes": list(raw.get("notes") or []),
        "error": str(raw.get("error") or ""),
        "environment": environment,
        "verification": "real_word_render",
        "secrets_included": False,
    }


def _failure_report(
    *,
    working: Path,
    pdf_path: Path | None,
    environment: Mapping[str, Any],
    error: str,
) -> dict[str, Any]:
    if pdf_path is not None:
        target = Path(pdf_path)
        if target.exists() and target.stat().st_size == 0:
            target.unlink(missing_ok=True)
    return {
        "status": "FAIL",
        "executed": True,
        "fields_updated": False,
        "toc_updated": False,
        "pagination_stable": False,
        "pdf_exported": False,
        "docx_saved": working.is_file(),
        "working_docx": str(working).replace("\\", "/"),
        "pdf_path": "",
        "environment": dict(environment),
        "error": error,
        "quit_called": False,
        "verification": "real_word_render",
        "secrets_included": False,
    }


def _probe_backend(
    backend: WordBackend,
    *,
    probe_factory: Callable[[Path], None] | None,
) -> dict[str, Any]:
    from docx import Document

    from app.word_renderer.oxml import add_page_field, add_toc_field

    with tempfile.TemporaryDirectory(prefix="word_finalizer_probe_") as raw_dir:
        directory = Path(raw_dir)
        source = directory / "probe.docx"
        working = directory / "probe_working.docx"
        pdf = directory / "probe.pdf"
        if probe_factory is not None:
            probe_factory(source)
        else:
            doc = Document()
            doc.add_paragraph("Probe Title")
            add_toc_field(doc.add_paragraph(), include_sections=False)
            add_page_field(doc.add_paragraph())
            doc.save(source)
        result = finalize_word_document(
            source_docx=source,
            working_docx=working,
            pdf_path=pdf,
            backend=backend,
            max_iterations=2,
        )
        capabilities = {
            "open_save": "PASS" if result.get("docx_saved") else "FAIL",
            "update_fields": "PASS" if result.get("fields_updated") else "FAIL",
            "update_toc": "PASS" if result.get("toc_updated") else "FAIL",
            "export_pdf": "PASS" if result.get("pdf_exported") else "FAIL",
        }
        return {
            "ok": result.get("status") in {"PASS", "PARTIAL"} and result.get("docx_saved"),
            "capabilities": capabilities,
            "error": result.get("error") or "",
            "status": result.get("status"),
        }


_POWERSHELL_SCRIPT = r"""
param(
    [Parameter(Mandatory = $true)][string]$InputDocx,
    [string]$OutputPdf = "",
    [Parameter(Mandatory = $true)][string]$ReportJson,
    [int]$MaxIterations = 3
)

$ErrorActionPreference = "Stop"
$report = [ordered]@{
    ok = $false
    status = "FAIL"
    backend = "powershell_com"
    backend_ready = $true
    fields_updated = $false
    toc_updated = $false
    pagination_stable = $false
    stability_iterations = 0
    page_count = 0
    toc_entries = @()
    chapter_starts = @()
    pdf_exported = $false
    docx_saved = $false
    word_was_already_running = $false
    created_new_word_instance = $false
    quit_called = $false
    documents_closed_by_us = 0
    error = ""
    notes = @()
}

function Write-Report {
    ($report | ConvertTo-Json -Compress -Depth 8) | Set-Content -LiteralPath $ReportJson -Encoding utf8
}

function Update-AllFields($document) {
    $document.Fields.Update() | Out-Null
    foreach ($section in $document.Sections) {
        foreach ($kind in 1, 2, 3) {
            try { $section.Headers.Item($kind).Range.Fields.Update() | Out-Null } catch {}
            try { $section.Footers.Item($kind).Range.Fields.Update() | Out-Null } catch {}
        }
    }
}

function Read-Toc($document) {
    $entries = @()
    if ($document.TablesOfContents.Count -lt 1) { return $entries }
    $toc = $document.TablesOfContents.Item(1)
    foreach ($para in $toc.Range.Paragraphs) {
        $text = ($para.Range.Text -replace "[\r\a]", "").Trim()
        if (-not $text) { continue }
        if ($text -match '^(.*?)[\.\s\t]+(\d+)$') {
            $entries += [ordered]@{ title = $Matches[1].Trim(); page = [int]$Matches[2]; raw = $text }
        } else {
            $entries += [ordered]@{ title = $text; page = $null; raw = $text }
        }
    }
    return $entries
}

$word = $null
$doc = $null
$createdNew = $false

try {
    $beforePids = @(Get-Process -Name WINWORD -ErrorAction SilentlyContinue | ForEach-Object { $_.Id })
    $report.word_was_already_running = ($beforePids.Count -gt 0)
    $word = New-Object -ComObject Word.Application
    Start-Sleep -Milliseconds 400
    $afterPids = @(Get-Process -Name WINWORD -ErrorAction SilentlyContinue | ForEach-Object { $_.Id })
    $newPids = @($afterPids | Where-Object { $beforePids -notcontains $_ })
    $createdNew = ($newPids.Count -gt 0)
    $report.created_new_word_instance = $createdNew

    $word.Visible = $false
    $word.DisplayAlerts = 0
    try { $word.ScreenUpdating = $false } catch {}
    try { $word.AutomationSecurity = 3 } catch { $report.notes += "AutomationSecurity could not be forced." }

    $doc = $word.Documents.Open($InputDocx, $false, $false, $false)

    $prevPages = -1
    $prevToc = ""
    $stable = $false
    $iterations = 0
    for ($i = 1; $i -le $MaxIterations; $i++) {
        $iterations = $i
        Update-AllFields $doc
        $report.fields_updated = $true
        $doc.Repaginate()
        if ($doc.TablesOfContents.Count -gt 0) {
            $doc.TablesOfContents.Item(1).Update()
            try { $doc.TablesOfContents.Item(1).UpdatePageNumbers() } catch {}
            $report.toc_updated = $true
        }
        $doc.Repaginate()
        $pages = [int]$doc.ComputeStatistics(2)
        $tocNow = ((Read-Toc $doc) | ConvertTo-Json -Compress -Depth 6)
        if (($i -gt 1) -and ($pages -eq $prevPages) -and ($tocNow -eq $prevToc)) {
            $stable = $true
            $report.page_count = $pages
            break
        }
        $prevPages = $pages
        $prevToc = $tocNow
        $report.page_count = $pages
    }
    $report.pagination_stable = [bool]$stable
    $report.stability_iterations = $iterations
    $report.toc_entries = @(Read-Toc $doc)

    $chapters = @()
    foreach ($para in $doc.Paragraphs) {
        $style = $null
        try { $style = [string]$para.Style.NameLocal } catch { continue }
        if ($style -eq "BookChapterTitle") {
            $title = ($para.Range.Text -replace "[\r\a]", "").Trim()
            $page = [int]$para.Range.Information(3)
            $adjusted = $page
            try { $adjusted = [int]$para.Range.Information(1) } catch {}
            $chapters += [ordered]@{
                title = $title
                page = $page
                page_adjusted = $adjusted
                odd = (($page % 2) -eq 1)
            }
        }
    }
    $report.chapter_starts = $chapters

    $doc.Save()
    $report.docx_saved = (Test-Path -LiteralPath $InputDocx)

    if ($OutputPdf) {
        $parent = Split-Path -Parent $OutputPdf
        if ($parent -and -not (Test-Path -LiteralPath $parent)) {
            New-Item -ItemType Directory -Path $parent | Out-Null
        }
        $wdExportFormatPDF = 17
        $doc.ExportAsFixedFormat($OutputPdf, $wdExportFormatPDF, $false, 0)
        $report.pdf_exported = ((Test-Path -LiteralPath $OutputPdf) -and ((Get-Item -LiteralPath $OutputPdf).Length -gt 0))
    }

    $report.ok = [bool]$report.docx_saved -and [bool]$report.fields_updated
    if ($report.ok -and $report.toc_updated -and $report.pagination_stable -and ((-not $OutputPdf) -or $report.pdf_exported)) {
        $report.status = "PASS"
    } elseif ($report.ok) {
        $report.status = "PARTIAL"
    } else {
        $report.status = "FAIL"
    }
}
catch {
    $report.error = [string]$_.Exception.Message
    $report.status = "FAIL"
    $report.ok = $false
}
finally {
    if ($null -ne $doc) {
        try { $doc.Close($false) | Out-Null; $report.documents_closed_by_us = 1 } catch {}
        try { [System.Runtime.InteropServices.Marshal]::ReleaseComObject($doc) | Out-Null } catch {}
    }
    if ($null -ne $word) {
        if ($createdNew) {
            try { $word.Quit() | Out-Null; $report.quit_called = $true } catch {}
        } else {
            $report.notes += "Existing Word session left running; only the document opened by this process was closed."
            $report.quit_called = $false
        }
        try { [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null } catch {}
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
    Write-Report
}

if (-not $report.ok) { exit 1 }
exit 0
"""


class PowerShellWordBackend:
    name = "powershell_com"

    def detect(self) -> dict[str, Any]:
        ready = word_progid_present() and powershell_available()
        return {
            "backend_ready": ready,
            "backend": self.name,
            "error": "" if ready else "PowerShell or Word.Application is unavailable.",
        }

    def finalize(
        self,
        *,
        working_docx: Path,
        pdf_path: Path | None,
        max_iterations: int,
    ) -> dict[str, Any]:
        if not powershell_available():
            raise WordFinalizerError("powershell.exe is not available.")
        with tempfile.TemporaryDirectory(prefix="word_finalizer_ps_") as raw_dir:
            directory = Path(raw_dir)
            script = directory / "finalize.ps1"
            report = directory / "report.json"
            script.write_text(_POWERSHELL_SCRIPT.lstrip("\n"), encoding="utf-8")
            command = [
                "powershell.exe",
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script),
                "-InputDocx",
                str(Path(working_docx).resolve()),
                "-ReportJson",
                str(report),
                "-MaxIterations",
                str(int(max_iterations)),
            ]
            if pdf_path is not None:
                command.extend(["-OutputPdf", str(Path(pdf_path).resolve())])
            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                check=False,
                timeout=900,
            )
            payload = _read_json_file(report)
            if not payload:
                raise WordFinalizerError(
                    "Word PowerShell finalizer produced no report: "
                    f"{(completed.stderr or completed.stdout or '').strip()[:800]}"
                )
            if completed.returncode != 0 and not payload.get("error"):
                payload["error"] = (
                    (completed.stderr or completed.stdout or "Word finalizer failed").strip()[:800]
                )
            return payload


class Win32WordBackend:
    name = "win32com"

    def detect(self) -> dict[str, Any]:
        return {
            "backend_ready": win32com_available() and word_progid_present(),
            "backend": self.name,
            "error": "" if win32com_available() else "win32com is not installed.",
        }

    def finalize(
        self,
        *,
        working_docx: Path,
        pdf_path: Path | None,
        max_iterations: int,
    ) -> dict[str, Any]:
        raise WordFinalizerError(
            "win32com is reserved as a detection backend only on this machine; "
            "use the PowerShell COM finalizer."
        )


class UnavailableWordBackend:
    name = "unavailable"

    def detect(self) -> dict[str, Any]:
        return {
            "backend_ready": False,
            "backend": self.name,
            "error": "Microsoft Word is not available.",
        }

    def finalize(
        self,
        *,
        working_docx: Path,
        pdf_path: Path | None,
        max_iterations: int,
    ) -> dict[str, Any]:
        raise WordFinalizerError("Microsoft Word is not available.")


class FakeWordBackend:
    """Test double. Never talks to Microsoft Word."""

    name = "fake"

    def __init__(
        self,
        *,
        available: bool = True,
        fail: bool = False,
        error: str = "injected COM failure",
        page_count: int = 12,
        toc_entries: list[dict[str, Any]] | None = None,
        chapter_starts: list[dict[str, Any]] | None = None,
        export_pdf: bool = True,
        pagination_stable: bool = True,
    ) -> None:
        self.available = available
        self.fail = fail
        self.error = error
        self.page_count = page_count
        self.toc_entries = toc_entries or [
            {"title": "Opening the Door", "page": 3, "raw": "Opening the Door\t3"}
        ]
        self.chapter_starts = chapter_starts or [
            {"title": "Opening the Door", "page": 3, "odd": True}
        ]
        self.export_pdf = export_pdf
        self.pagination_stable = pagination_stable

    def detect(self) -> dict[str, Any]:
        return {
            "backend_ready": self.available,
            "backend": self.name,
            "error": "" if self.available else "Microsoft Word is not available.",
        }

    def finalize(
        self,
        *,
        working_docx: Path,
        pdf_path: Path | None,
        max_iterations: int,
    ) -> dict[str, Any]:
        del max_iterations
        if self.fail:
            raise WordFinalizerError(self.error)
        if pdf_path is not None and self.export_pdf:
            Path(pdf_path).parent.mkdir(parents=True, exist_ok=True)
            Path(pdf_path).write_bytes(b"%PDF-1.4\n% fake probe\n")
        return {
            "ok": True,
            "status": "PASS",
            "fields_updated": True,
            "toc_updated": True,
            "pagination_stable": self.pagination_stable,
            "stability_iterations": 2,
            "page_count": self.page_count,
            "toc_entries": self.toc_entries,
            "chapter_starts": self.chapter_starts,
            "pdf_exported": bool(pdf_path is not None and self.export_pdf),
            "docx_saved": Path(working_docx).is_file(),
            "word_was_already_running": False,
            "created_new_word_instance": True,
            "quit_called": True,
            "documents_closed_by_us": 1,
            "error": "",
            "notes": ["fake backend"],
        }


def _read_json_file(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    raw = path.read_text(encoding="utf-8-sig")
    if not raw.strip():
        return {}
    payload = json.loads(raw)
    return payload if isinstance(payload, dict) else {}


__all__ = [
    "FakeWordBackend",
    "PowerShellWordBackend",
    "UnavailableWordBackend",
    "Win32WordBackend",
    "WordBackend",
    "WordFinalizerError",
    "default_backend",
    "detect_word_environment",
    "finalize_word_document",
    "finalizer_readiness",
    "win32com_available",
    "word_progid_present",
]
